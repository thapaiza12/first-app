import unittest,sys,tempfile,base64,io,zipfile
from pathlib import Path
from PIL import Image
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'server'))
import documents
from unittest.mock import patch
class DocumentReading(unittest.TestCase):
 def test_word_textboxes_and_headers_are_read_including_contact_and_history(self):
  xml='''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>氏名：テスト 太郎</w:t></w:r></w:p><w:p><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>現住所：東京都千代田区1-1</w:t></w:r></w:p><w:p><w:r><w:t>電話番号：090-1234-5678</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>2020年4月 株式会社テスト 入社</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>'''
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'fixture.docx'
   with zipfile.ZipFile(p,'w') as z:
    z.writestr('word/document.xml',xml);z.writestr('word/header1.xml','<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>メール：test@example.com</w:t></w:r></w:p></w:hdr>')
   text=documents.extract(str(p),p.name)['text']
   for value in ['東京都千代田区1-1','090-1234-5678','株式会社テスト','test@example.com']:self.assertIn(value,text)
 def test_pdf_keeps_both_pages_and_sends_legible_full_page_images(self):
  import fitz
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'cv.pdf';pdf=fitz.open()
   for content in ['Address: Tokyo 1-1; Phone: 090-1234-5678','Work history: Sample Company 2020-04; Education: Sample University']:
    page=pdf.new_page(width=595,height=842);page.insert_text((30,700),content,fontsize=9)
   pdf.save(p);pdf.close();result=documents.extract(str(p),p.name)
   self.assertIn('090-1234-5678',result['text']);self.assertIn('Sample Company',result['text']);self.assertEqual(len(result['pages']),2)
   for page in result['pages']:
    image=Image.open(io.BytesIO(base64.b64decode(page.split(',')[1])))
    self.assertGreaterEqual(image.width,1400)
 def test_native_pdf_fast_path_does_not_skip_image_based_sections(self):
  import fitz
  with tempfile.TemporaryDirectory() as d:
   for has_image in [False,True]:
    p=Path(d)/('mixed.pdf' if has_image else 'native.pdf');pdf=fitz.open();page=pdf.new_page()
    page.insert_text((30,50),'Native CV text and contact details. '*5,fontsize=8)
    if has_image:
     buffer=io.BytesIO();Image.new('RGB',(400,600),'white').save(buffer,format='PNG');page.insert_image(fitz.Rect(30,150,500,700),stream=buffer.getvalue())
    pdf.save(p);pdf.close();result=documents.extract(str(p),p.name)
    self.assertEqual(result['textReady'],not has_image)
 def test_scanned_images_keep_readable_resolution_for_gemini(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'scan.png';Image.new('RGB',(1800,2600),'white').save(p)
   with patch.object(documents,'ocr_image',return_value='氏名：テスト'):
    result=documents.extract(str(p),p.name)
   image=Image.open(io.BytesIO(base64.b64decode(result['pages'][0].split(',')[1])))
   self.assertEqual(image.size,(1800,2600))

class FormatCoverage(unittest.TestCase):
 def test_text_word_excel_and_powerpoint_tables_have_complete_native_text(self):
  from docx import Document
  from openpyxl import Workbook
  from pptx import Presentation
  from pptx.util import Inches
  with tempfile.TemporaryDirectory() as d:
   root=Path(d)
   for suffix in ['txt','csv','md','json']:
    p=root/('job.'+suffix);p.write_text('求人票\n輸出入業務と英語力\n最終行',encoding='utf-8')
    result=documents.extract(str(p),p.name);self.assertTrue(result['textReady']);self.assertIn('最終行',result['text'])
   p=root/'job.docx';doc=Document();doc.add_paragraph('求人票：輸出入業務');doc.add_table(rows=1,cols=1).cell(0,0).text='最終セル';doc.save(p)
   result=documents.extract(str(p),p.name);self.assertTrue(result['textReady']);self.assertIn('最終セル',result['text'])
   p=root/'job.xlsx';book=Workbook();book.active['A1']='第一シート';book.create_sheet('最終シート')['C3']='最終セル';book.save(p)
   result=documents.extract(str(p),p.name);self.assertTrue(result['textReady']);self.assertIn('最終シート',result['text']);self.assertIn('最終セル',result['text'])
   p=root/'job.pptx';deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_table(1,1,0,0, Inches(4), Inches(1)).table.cell(0,0).text='表内の応募条件';deck.save(p)
   result=documents.extract(str(p),p.name);self.assertTrue(result['textReady']);self.assertIn('表内の応募条件',result['text'])
 def test_jpeg_png_and_multi_page_tiff_are_ocr_ready_only_if_every_image_was_read(self):
  with tempfile.TemporaryDirectory() as d:
   for suffix in ['jpg','png','tiff']:
    p=Path(d)/('scan.'+suffix);image=Image.new('RGB',(300,400),'white')
    image.save(p,save_all=suffix=='tiff',append_images=[image] if suffix=='tiff' else [])
    expected=2 if suffix=='tiff' else 1
    with patch.object(documents,'ocr_image',side_effect=['第'+str(i)+'ページの原文' for i in range(expected)]):result=documents.extract(str(p),p.name)
    self.assertFalse(result['textReady']);self.assertTrue(result['ocrTextReady']);self.assertEqual(result['pageCount'],expected)
    self.assertIn('第'+str(expected-1)+'ページの原文',result['text'])
   with patch.object(documents,'ocr_image',side_effect=['最初のページ','']):result=documents.extract(str(p),p.name)
   self.assertFalse(result['ocrTextReady'])
 def test_mac_ocr_uses_native_vision_and_removes_its_temporary_image(self):
  from types import SimpleNamespace
  buffer=io.BytesIO();Image.new('RGB',(300,400),'white').save(buffer,format='PNG');paths=[]
  def invoke(args,**kwargs):
   paths.append(Path(args[-1]));self.assertTrue(paths[-1].exists());self.assertIn('mac_ocr.js',args[-2]);return SimpleNamespace(returncode=0,stdout='{"text":"画像の日本語"}'.encode())
  with patch.object(documents.sys,'platform','darwin'),patch.object(documents.subprocess,'run',side_effect=invoke):self.assertEqual(documents.ocr_image(buffer.getvalue()),'画像の日本語')
  self.assertFalse(paths[0].exists())
 def test_rtf_is_decoded_as_text_instead_of_returning_control_codes(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'job.rtf';p.write_text(r'{\rtf1\ansi Job offer\par Final line}')
   result=documents.extract(str(p),p.name);self.assertTrue(result['textReady']);self.assertIn('Final line',result['text']);self.assertNotIn('\\rtf1',result['text'])

if __name__=='__main__':unittest.main()
