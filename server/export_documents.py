"""Distinct Japanese résumé and career-history layouts; no fabricated personal fields."""
import base64,io
from datetime import date,datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape

TITLE='学歴・職歴（各別にまとめて書く）'
MOTIVATION='志望の動機、特技、好きな学科、アピールポイントなど'
PREFERENCES='本人希望記入欄（特に給料・職種・勤務時間・勤務地・その他についての希望などがあれば記入）'

def asof(d):
    try:return date.fromisoformat(d.get('documentDate',''))
    except ValueError:return datetime.now(ZoneInfo('Asia/Tokyo')).date()

def jpdate(value):
    if not value:return ''
    try:
        parts=value.split('-')
        return f'{int(parts[0])}年{int(parts[1])}月'+(f'{int(parts[2])}日' if len(parts)>2 else '')
    except (ValueError,IndexError):return value

def birth(d):
    raw=d.get('birthDate','')
    try:
        b=date.fromisoformat(raw);today=asof(d);age=today.year-b.year-((today.month,today.day)<(b.month,b.day))
        return f'{jpdate(raw)}生（満{age}歳）' if age>=0 else jpdate(raw)
    except ValueError:return raw

def period(row):
    start=jpdate(row.get('date',''));end=jpdate(row.get('end',''))
    if row.get('current'):end='現在'
    return start+(' 〜 '+end if end else '')

def timeline(d):
    out=[]
    for label,key in [('学歴','education'),('職歴','experience')]:
        out.append(['','',label])
        for r in d.get(key,[]):
            raw=r.get('date','').split('-');year=raw[0] if raw else '';month=raw[1] if len(raw)>1 else ''
            value='\n'.join(filter(None,[r.get('name',''),r.get('detail','')]))
            if r.get('end') or r.get('current'):value+='　（'+period(r)+'）'
            out.append([year,month,value])
    out.append(['','','以上'])
    return out

def qualification_rows(d):
    rows=[]
    for r in d.get('licenses',[]):
        raw=r.get('date','').split('-');rows.append([raw[0] if raw else '',raw[1] if len(raw)>1 else '',r.get('name','')])
    return rows

def career_tail(d):
    qualifications='資格\n'+'\n'.join(r.get('name','')+'　'+jpdate(r.get('date','')) for r in d.get('licenses',[]))
    languages='語学力\n'+'\n'.join('　'.join(filter(None,[r.get('language',''),r.get('level',''),r.get('qualification',''),jpdate(r.get('date',''))])) for r in d.get('languages',[]))
    pr=d.get('additional','')
    if d.get('motivation') and d['motivation'] not in pr:pr+='\n\n'+d['motivation']
    return [('4. テクニカルスキル',d.get('techSkills','')),('5. 資格・語学力',qualifications+'\n\n'+languages),('6. 自己PR・志望動機',pr.strip())]

def render_document(d,fmt,out):
    if d.get('type')!='shokumukeirekisho':
        import resume_layout
        return resume_layout.word(d,out) if fmt=='docx' else resume_layout.pdf(d,out)
    if fmt=='docx':return word(d,out)
    return pdf(d,out)

def pdf(d,out):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,LongTable,TableStyle,Image,PageBreak,KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER,TA_RIGHT
    font=Path(__file__).parent/'fonts/NotoSerifJP.ttf'
    if 'JP' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('JP',str(font)))
    width=493
    normal=ParagraphStyle('jp',fontName='JP',fontSize=10,leading=16,wordWrap='CJK')
    small=ParagraphStyle('small',parent=normal,fontSize=8,leading=12)
    name_style=ParagraphStyle('name',parent=normal,fontSize=18,leading=26)
    heading=ParagraphStyle('head',parent=normal,fontSize=12,leading=21,spaceBefore=13,spaceAfter=7,keepWithNext=True)
    title=ParagraphStyle('title',parent=normal,fontSize=22,leading=32,spaceAfter=10,alignment=TA_CENTER)
    right=ParagraphStyle('right',parent=small,alignment=TA_RIGHT)
    p=lambda value,style=normal:Paragraph(escape(str(value or '　')).replace('\n','<br/>'),style)
    story=[]
    def grid(rows,widths,heights=None,repeat=0):
        t=LongTable(rows,colWidths=widths,minRowHeights=heights,repeatRows=repeat,splitInRow=1)
        t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.black),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        return t
    def chronology(rows,label,min_rows):
        padded=rows+[['','','']] * max(0,min_rows-len(rows))
        return grid([[p('年',small),p('月',small),p(label,small)]]+[[p(v,small) for v in row] for row in padded],[43,28,width-71],[22]*(len(padded)+1),1)
    def box(label,value,height):
        return grid([[p(label,small)],[p(value)]],[width],[22,height])
    resume=d.get('type')!='shokumukeirekisho'
    if resume:
        story.extend([p('履　歴　書',title),p(jpdate(asof(d).isoformat())+' 現在',right),Spacer(1,8)])
        photo=p('写真を貼る位置\n縦40mm × 横30mm',small)
        if d.get('photo'):
            image=base64.b64decode(d['photo'].split(',',1)[1]);photo=Image(io.BytesIO(image),width=85,height=113,kind='proportional')
        photo_frame=Table([[photo]],colWidths=[91],rowHeights=[119]);photo_frame.setStyle(TableStyle([('BOX',(0,0),(-1,-1),.4,colors.grey),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('ALIGN',(0,0),(-1,-1),'CENTER'),('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]))
        personal=grid([[p('ふりがな　'+d.get('furigana',''),small)],[[p('氏名',small),p(d.get('name',''),name_style)]],[p(birth(d)+'　　性別（任意） '+d.get('gender',''),small)]],[width-105],[23,51,30])
        outer=Table([[personal,photo_frame]],colWidths=[width-99,99]);outer.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),0),('BOTTOMPADDING',(0,0),(-1,-1),0)]));story.append(outer)
        story.append(Spacer(1,10))
        contact=[[p('ふりがな　'+d.get('addressKana',''),small),p('電話　'+d.get('phone',''),small)],[p('現住所　〒'+d.get('postalCode','')+'\n'+d.get('address','')),p('E-mail\n'+d.get('email',''),small)],[p('ふりがな　'+d.get('alternateAddressKana',''),small),p('電話　'+d.get('alternatePhone',''),small)],[p('連絡先　〒'+d.get('alternatePostalCode','')+'\n'+d.get('alternateAddress',''),small),p('（現住所以外に連絡を希望する場合のみ記入）',small)]]
        story.extend([grid(contact,[width-135,135],[20,44,20,42]),Spacer(1,12)])
        rows=timeline(d)
        story.append(chronology(rows[:14],TITLE,14))
        story.append(PageBreak())
        story.extend([chronology(rows[14:],TITLE,5),Spacer(1,12),chronology(qualification_rows(d),'免許・資格',5),Spacer(1,12)])
        story.append(box(MOTIVATION,'\n\n'.join(filter(None,[d.get('motivation',''),d.get('strengths','')])),105))
        story.extend([Spacer(1,12),box(PREFERENCES,d.get('preferences',''),85)])
    else:
        story.append(p('職 務 経 歴 書',title))
        # A career-history document uses narrative headings, not résumé personal-data boxes.
        for label,key in [('1. 職務要約','summary'),('2. 活かせる経験・知識・技術','skills')]:
            story.append(p(label,heading))
            if key=='summary':story.extend([p('氏名：'+d.get('name','')),p('記入日：'+jpdate(asof(d).isoformat())+' 現在'),Spacer(1,7)])
            story.append(p(d.get(key,'')))
        story.append(p('3. 職務経歴',heading))
        if not d.get('experience'):story.append(p(''))
        for i,r in enumerate(d.get('experience',[]),1):
            story.append(p(str(i)+'. '+r.get('name',''),heading))
            meta=[]
            for label,key in [('在籍期間',None),('事業内容','business'),('雇用形態','employmentType'),('役職・担当','position')]:
                value=period(r) if key is None else r.get(key,'')
                if value:meta.append([p(label,small),p(value)])
            if meta:story.append(grid(meta,[85,width-85]))
            story.extend([Spacer(1,7),p('担当業務・実績',small)])
            for line in (r.get('detail','') or '　').split('\n'):story.append(p(line))
            story.append(Spacer(1,10))
        for label,value in career_tail(d):
            story.append(p(label,heading))
            for line in (value or '　').split('\n'):story.append(p(line))
        story.extend([Spacer(1,16),p('以上',right)])
    def footer(canvas,doc):
        canvas.setFont('JP',8);canvas.drawRightString(545,25,str(doc.page))
    SimpleDocTemplate(out,pagesize=(595.28,841.89),leftMargin=51,rightMargin=51,topMargin=37,bottomMargin=38,title='履歴書' if resume else '職務経歴書',author=d.get('name','')).build(story,onFirstPage=footer,onLaterPages=footer)

def word(d,out):
    from docx import Document
    from docx.shared import Mm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    doc=Document();sec=doc.sections[0];sec.page_height=Mm(297);sec.page_width=Mm(210)
    sec.top_margin=sec.bottom_margin=Mm(13);sec.left_margin=sec.right_margin=Mm(18)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.name='Noto Serif JP';style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif JP')
        if name=='Normal':style.font.size=Pt(10);style.paragraph_format.space_after=Pt(4)
    def para(value='',align=None):
        p=doc.add_paragraph(value)
        if align is not None:p.alignment=align
        return p
    def table(rows,widths=None):
        t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for values in rows:
            cells=t.add_row().cells
            for i,value in enumerate(values):cells[i].text=str(value or '　')
        if widths:
            for r in t.rows:
                for c,w in zip(r.cells,widths):c.width=Mm(w)
        return t
    def chron(rows,label,count):
        padded=rows+[['','','']]*max(0,count-len(rows))
        return table([['年','月',label]]+padded,[15,10,149])
    resume=d.get('type')!='shokumukeirekisho'
    title=doc.add_paragraph('履 歴 書' if resume else '職 務 経 歴 書','Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if resume:para(jpdate(asof(d).isoformat())+' 現在',WD_ALIGN_PARAGRAPH.RIGHT)
    if resume:
        t=table([['ふりがな　'+d.get('furigana',''),''],['氏名　'+d.get('name',''),''],[birth(d)+'　性別（任意） '+d.get('gender',''),'']],[140,34])
        t.cell(1,0).paragraphs[0].runs[0].font.size=Pt(18)
        photo_cell=t.cell(0,1).merge(t.cell(2,1))
        if d.get('photo'):photo_cell.paragraphs[0].add_run().add_picture(io.BytesIO(base64.b64decode(d['photo'].split(',',1)[1])),width=Mm(30))
        else:photo_cell.text='写真を貼る位置\n縦40mm × 横30mm'
        table([['ふりがな　'+d.get('addressKana',''),'電話　'+d.get('phone','')],['現住所　〒'+d.get('postalCode','')+'\n'+d.get('address',''),'E-mail　'+d.get('email','')],['ふりがな　'+d.get('alternateAddressKana',''),'電話　'+d.get('alternatePhone','')],['連絡先　〒'+d.get('alternatePostalCode','')+'\n'+d.get('alternateAddress',''),'（現住所以外に連絡を希望する場合のみ記入）']],[124,50])
        para();rows=timeline(d);chron(rows[:14],TITLE,14);doc.add_page_break();chron(rows[14:],TITLE,5);para();chron(qualification_rows(d),'免許・資格',5);para()
        table([[MOTIVATION],['\n\n'.join(filter(None,[d.get('motivation',''),d.get('strengths','')])) or '\n\n']]);para();table([[PREFERENCES],[d.get('preferences','') or '\n\n']])
    else:
        for label,key in [('1. 職務要約','summary'),('2. 活かせる経験・知識・技術','skills')]:
            doc.add_heading(label,1)
            if key=='summary':para('氏名：'+d.get('name',''));para('記入日：'+jpdate(asof(d).isoformat())+' 現在')
            para(d.get(key,''))
        doc.add_heading('3. 職務経歴',1)
        for i,r in enumerate(d.get('experience',[]),1):
            doc.add_heading(str(i)+'. '+r.get('name',''),2)
            rows=[]
            for label,key in [('在籍期間',None),('事業内容','business'),('雇用形態','employmentType'),('役職・担当','position')]:
                value=period(r) if key is None else r.get(key,'')
                if value:rows.append([label,value])
            if rows:table(rows,[30,144])
            doc.add_heading('担当業務・実績',3);para(r.get('detail',''))
        for label,value in career_tail(d):
            doc.add_heading(label,1);para(value)
        para('以上',WD_ALIGN_PARAGRAPH.RIGHT)
    doc.save(out)
