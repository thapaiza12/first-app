"""Measured, fixed two-sheet résumé layout shared by PDF and editable Word output.
Content that cannot fit at a readable size is rejected, never clipped or put on page 3.
"""
from pathlib import Path
import base64,io,re
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from export_documents import asof,jpdate,birth,TITLE,MOTIVATION,PREFERENCES

PAGE_W,PAGE_H=595.28,841.89
X,W=42,511.28

def init_font():
    if 'JP' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('JP',str(Path(__file__).parent/'fonts/NotoSerifJP.ttf')))

def lines(value,size,width):
    init_font();result=[]
    for paragraph in str(value or '').split('\n'):
        current=''
        for char in paragraph:
            if current and pdfmetrics.stringWidth(current+char,'JP',size)>width:
                result.append(current);current=char
            else:current+=char
        result.append(current)
    return result or ['']

def fit(value,width,height,label,base=10.5,minimum=9):
    for size in [base-i*.5 for i in range(int((base-minimum)*2)+1)]:
        wrapped=lines(value,size,width*.91)
        if len(wrapped)*size*1.35<=height-8:
            return {'text':str(value or ''),'lines':wrapped,'size':size,'leading':size*1.35}
    raise ValueError(f'履歴書を2ページに収められません。「{label}」の内容を短くしてください。内容は保存されたままです。')

def normalize_jobs(rows):
    result=[]
    for original in rows:
        r=dict(original);m=re.match(r'^(.*?)[\s　]*(入社|退社|退職)\s*$',r.get('name',''));base=m[1].strip() if m else r.get('name','')
        if m and m[2] in ['退社','退職']:
            if result and result[-1].get('name')==base and not result[-1].get('end'):
                result[-1]['end']=r.get('date','');result[-1]['current']=False;continue
            r.update(date='',end=r.get('date',''),current=False)
        r['name']=base;result.append(r)
    return result

def employment_events(rows):
    out=[]
    for r in normalize_jobs(rows):
        out.append({'date':r.get('date',''),'name':r.get('name','')+'　入社'})
        out.append({'date':'' if r.get('current') else r.get('end',''),'name':'現在に至る' if r.get('current') else r.get('name','')+'　退社'})
    return out

def education_events(rows):
    out=[]
    event=re.compile(r'(入学|卒業|修了|中退|退学|在学中)\s*$')
    for original in rows:
        r=dict(original);name=r.get('name','').strip();detail=r.get('detail','').strip()
        parts=[p for p in [*name.splitlines(),*detail.splitlines()] if p]
        explicit=list(dict.fromkeys(p for p in parts if event.search(p)))
        if explicit:
            extra=[p for p in parts if not event.search(p) and not any(e.startswith(p) for e in explicit)]
            for i,p in enumerate(explicit):out.append({**r,'date':r.get('end','') if i else r.get('date',''),'end':'','name':p,'detail':'' if i else '\n'.join(extra)})
        elif r.get('end'):
            out.extend([{**r,'name':name+'　入学','end':'','detail':''},{**r,'date':r['end'],'end':'','name':name+'　卒業','detail':detail}])
        else:out.append({**r,'name':name,'detail':detail})
    return out

def chronology(d):
    result=[['','','学歴']]
    for r in education_events(d.get('education',[])):
        date=r.get('date','').split('-');result.append([date[0] if date else '',date[1] if len(date)>1 else '',r.get('name','')+ ('\n'+r['detail'] if r.get('detail') else '')])
    result.append(['','',''])
    result.append(['','','職歴'])
    for r in employment_events(d.get('experience',[])):
        date=r.get('date','').split('-');result.append([date[0] if date else '',date[1] if len(date)>1 else '',r.get('name','')])
    result.append(['','','以上']);return result

def measured_rows(rows,size,widths):
    result=[]
    for row in rows:
        cells=[{'text':str(v or ''),'lines':lines(v,size,(w-12)*.91),'size':size,'leading':size*1.35} for v,w in zip(row,widths)]
        if row[2] in ['学歴','職歴']:
            cells[2].update(size=size+1,leading=(size+1)*1.35,bold=True,align='center')
        elif row[2]=='以上':cells[2]['align']='right'
        height=max(20,max(len(c['lines'])*c['leading'] for c in cells)+8)
        result.append({'cells':cells,'height':height})
    return result

def pad_rows(rows,height,size,widths):
    rows=list(rows);remaining=height-sum(r['height'] for r in rows)
    while remaining>=20:
        rows.append({'cells':[{'text':'','lines':[''],'size':size,'leading':size*1.35} for _ in widths],'height':20});remaining-=20
    if remaining>0:
        if rows:rows[-1]['height']+=remaining
        else:rows=[{'cells':[{'text':'','lines':[''],'size':size,'leading':size*1.35} for _ in widths],'height':height}]
    return rows

def tables(d):
    widths=[43,28,W-71];all_rows=chronology(d)
    qualification=[]
    for r in d.get('licenses',[]):
        date=r.get('date','').split('-');qualification.append([date[0] if date else '',date[1] if len(date)>1 else '',r.get('name','')])
    motivation='\n\n'.join(filter(None,[d.get('motivation',''),d.get('strengths','')]))
    for size in [10.5,10,9.5,9]:
        measured=measured_rows(all_rows,size,widths);left=[];right=[];used=0
        for r in measured:
            if not right and used+r['height']<=422:left.append(r);used+=r['height']
            else:right.append(r)
        # Page two shares one budget. Empty table rows and empty text areas do not
        # reserve space that another section needs. Content is never truncated.
        qrows=measured_rows(qualification,size,widths)
        right=right or pad_rows([],20,size,widths)
        qrows=qrows or pad_rows([],20,size,widths)
        bodies={}
        for key,value,minimum in [('motivation',motivation,40),('preferences',d.get('preferences',''),20)]:
            wrapped=lines(value,size,(W-12)*.91)
            bodies[key]=max(minimum,len(wrapped)*size*1.35+12)
        heights={'chronology':22+sum(r['height'] for r in right)+2,'qualifications':22+sum(r['height'] for r in qrows)+2,'motivation':22+bodies['motivation'],'preferences':30+bodies['preferences']}
        if sum(heights.values())+36<=716:
            return widths,[pad_rows(left,422,size,widths),right],qrows,{'heights':heights,'bodies':bodies,'size':size}
    raise ValueError('空白行・空欄の余白・文字サイズを自動調整しましたが、全内容を読みやすい2ページに収められません。入力内容は保存されています。詳しい職務内容は職務経歴書へ移してください。')

def contact_block(d):
    values=['Phone: '+d.get('phone',''),'Email: '+d.get('email','')]
    init_font()
    widest=max(pdfmetrics.stringWidth(v,'JP',9) for v in values)
    size=min(9,9*126/max(widest,1))
    if size<5.5:raise ValueError('メールアドレスが長すぎます。入力内容を確認してください。')
    return {'text':'\n'.join(values),'lines':values,'size':size,'leading':size*1.5}

def layout(d):
    init_font();widths,jobs,quals,plan=tables(d)
    # Every value is measured before either format is written.
    fields={
      'kana':fit('ふりがな　'+d.get('furigana',''),W-115,25,'ふりがな',9,9),
      'name':fit(d.get('name',''),W-115,52,'氏名',18,12),
      'birth':fit(birth(d)+'　性別（任意） '+d.get('gender',''),W-115,29,'生年月日・性別',9,9),
      'address':fit('ふりがな　'+d.get('addressKana','')+'\n現住所　〒'+d.get('postalCode','')+'\n'+d.get('address',''),W-150,48,'現住所',9,9),
      'contact':contact_block(d),
      'alternate':fit('ふりがな　'+d.get('alternateAddressKana','')+'\n連絡先　〒'+d.get('alternatePostalCode','')+'\n'+d.get('alternateAddress',''),W-150,48,'連絡先',9,9),
      'alternatePhone':fit('Phone: '+d.get('alternatePhone','')+'\n（現住所以外に連絡を希望する場合のみ記入）',138,48,'連絡先電話',8,8),
      'motivation':fit('\n\n'.join(filter(None,[d.get('motivation',''),d.get('strengths','')])),W-12,plan['bodies']['motivation'],'志望の動機・特技・自己PR',plan['size'],plan['size']),
      'preferences':fit(d.get('preferences',''),W-12,plan['bodies']['preferences'],'本人希望記入欄',plan['size'],plan['size']),
    }
    return {'pages':2,'widths':widths,'jobs':jobs,'qualifications':quals,'fields':fields,'page2':plan['heights']}

def pdf(d,out):
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    model=layout(d);c=canvas.Canvas(out,pagesize=(PAGE_W,PAGE_H));c.setTitle('履歴書');c.setAuthor(d.get('name',''))
    def rule(x,y,w,h):c.setLineWidth(.5);c.rect(x,PAGE_H-y-h,w,h,stroke=1,fill=0)
    def text(block,x,y):
        c.setFont('JP',block['size'])
        for i,line in enumerate(block['lines']):
            baseline=PAGE_H-y-block['size']-i*block['leading']
            if block.get('bold'):
                c.saveState();c.setLineWidth(.18);t=c.beginText(x,baseline);t.setFont('JP',block['size']);t.setTextRenderMode(2);t.textOut(line);c.drawText(t);c.restoreState()
            else:c.drawString(x,baseline,line)
    def label(value,x,y,size=8):text({'lines':[value],'size':size,'leading':size*1.35},x,y)
    def grid(rows,x,y,widths,label_value,total_height):
        rule(x,y,sum(widths),total_height)
        for offset in [widths[0],widths[0]+widths[1]]:c.line(x+offset,PAGE_H-y,x+offset,PAGE_H-y-total_height)
        for value,offset in [('年',0),('月',widths[0]),(label_value,widths[0]+widths[1])]:label(value,x+offset+5,y+5)
        cursor=y+22;c.line(x,PAGE_H-cursor,x+sum(widths),PAGE_H-cursor)
        for row in rows:
            offset=0
            for field,width in zip(row['cells'],widths):
                tx=x+offset+6
                if field.get('align')=='center':tx=x+offset+(width-pdfmetrics.stringWidth(field['lines'][0],'JP',field['size']))/2
                elif field.get('align')=='right':tx=x+offset+width-6-pdfmetrics.stringWidth(field['lines'][0],'JP',field['size'])
                text(field,tx,cursor+4);offset+=width
            cursor+=row['height'];c.line(x,PAGE_H-cursor,x+sum(widths),PAGE_H-cursor)
    def header_footer(page):
        c.setFont('JP',8);c.drawRightString(PAGE_W-42,24,str(page))
    label('履　歴　書',X,42,22);label(jpdate(asof(d).isoformat())+' 現在',X+W-145,61,9)
    main=W-100;rule(X,86,main,120);rule(X+main+7,89,85,113)
    text(model['fields']['kana'],X+6,90);c.line(X,PAGE_H-115,X+main,PAGE_H-115)
    label('氏名',X+6,119,8);text(model['fields']['name'],X+6,133);c.line(X,PAGE_H-177,X+main,PAGE_H-177)
    text(model['fields']['birth'],X+6,181)
    if d.get('photo'):
        image=ImageReader(io.BytesIO(base64.b64decode(d['photo'].split(',',1)[1])))
        c.drawImage(image,X+main+7,PAGE_H-202,width=85,height=113,preserveAspectRatio=True,anchor='c')
    else:label('写真 40mm × 30mm',X+main+10,120,7)
    for y,left,right in [(206,'address','contact'),(254,'alternate','alternatePhone')]:
        rule(X,y,W,48);c.line(X+W-144,PAGE_H-y,X+W-144,PAGE_H-y-48)
        text(model['fields'][left],X+6,y+3);text(model['fields'][right],X+W-138,y+3)
    grid(model['jobs'][0],X,318,model['widths'],TITLE,446);header_footer(1);c.showPage()
    heights=model['page2'];y=42
    grid(model['jobs'][1],X,y,model['widths'],TITLE,heights['chronology']);y+=heights['chronology']+12
    grid(model['qualifications'],X,y,model['widths'],'免許・資格',heights['qualifications']);y+=heights['qualifications']+12
    for key,heading,labelh in [('motivation',MOTIVATION,22),('preferences',PREFERENCES,30)]:
        h=heights[key];rule(X,y,W,h);c.line(X,PAGE_H-y-labelh,X+W,PAGE_H-y-labelh)
        heading_block=fit(heading,W-12,labelh,heading,8,8);text(heading_block,X+6,y+3);text(model['fields'][key],X+6,y+labelh+4)
        y+=h+12
    header_footer(2);c.showPage();c.save()
    import fitz
    check=fitz.open(out)
    if len(check)!=2:raise ValueError('2ページの履歴書を生成できませんでした。')
    check.close()

def word(d,out):
    from docx import Document
    from docx.shared import Pt,Mm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_BREAK,WD_ALIGN_PARAGRAPH
    model=layout(d);doc=Document();sec=doc.sections[0];sec.page_width=Pt(PAGE_W);sec.page_height=Pt(PAGE_H)
    sec.left_margin=sec.right_margin=Pt(42);sec.top_margin=sec.bottom_margin=Pt(42)
    style=doc.styles['Normal'];style.font.name='Noto Serif JP';style.font.size=Pt(9);style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif JP')
    style.paragraph_format.space_after=Pt(0);style.paragraph_format.space_before=Pt(0)
    def cell_format(cell):
        tcPr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
        for edge in ['top','left','bottom','right']:
            el=OxmlElement('w:'+edge);el.set(qn('w:w'),'0');el.set(qn('w:type'),'dxa');m.append(el)
        tcPr.append(m);no_wrap=OxmlElement('w:noWrap');tcPr.append(no_wrap)
    def write(cell,blocks,padding=3):
        cell_format(cell);cell.text=''
        for i,b in enumerate(blocks):
            paragraph=cell.paragraphs[0] if i==0 else cell.add_paragraph()
            paragraph.paragraph_format.line_spacing=Pt(b['leading']);paragraph.paragraph_format.left_indent=Pt(padding);paragraph.paragraph_format.right_indent=Pt(padding)
            paragraph.alignment={'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT}.get(b.get('align'),WD_ALIGN_PARAGRAPH.LEFT)
            run=paragraph.add_run('\n'.join(b['lines']));run.bold=b.get('bold',False);run.font.size=Pt(b['size']);run.font.name='Noto Serif JP';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif JP')
    def block(value,size=9):return {'lines':[value],'size':size,'leading':size*1.35}
    def fixrow(row,height):
        row.height=Pt(height);row.height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
    def nested(cell,rows,widths,heights):
        cell_format(cell);cell.text='';cell.paragraphs[0].paragraph_format.line_spacing=Pt(1);cell.paragraphs[0].runs
        t=cell.add_table(rows=0,cols=len(widths));t.autofit=False;t.style='Table Grid'
        for col,width in zip(t.columns,widths):col.width=Pt(width)
        for values,height in zip(rows,heights):
            row=t.add_row();fixrow(row,height)
            for target,field,width in zip(row.cells,values,widths):target.width=Pt(width);write(target,[field])
        # Word requires a final paragraph after nested tables. Keep it at one point.
        for p in [cell.paragraphs[0],cell.paragraphs[-1]]:
            p.paragraph_format.line_spacing=Pt(1);p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
            for run in p.runs:run.font.size=Pt(1)
        return t
    def outer(heights):
        t=doc.add_table(rows=len(heights),cols=1);t.autofit=False;t.columns[0].width=Pt(W)
        for row,h in zip(t.rows,heights):fixrow(row,h);cell_format(row.cells[0])
        return [r.cells[0] for r in t.rows]
    def chronology_cell(cell,rows,label):
        vals=[[block('年',8),block('月',8),block(label,8)]]+[r['cells'] for r in rows]
        heights=[22]+[r['height'] for r in rows];heights[-1]-=2
        nested(cell,vals,model['widths'],heights)
    first=outer([44,120,48,48,16,446]);write(first[0],[block('履 歴 書',20),block(jpdate(asof(d).isoformat())+' 現在',9)])
    t=nested(first[1],[[block(''),block('')]], [W-100,100],[118]);left,right=t.rows[0].cells
    write(left,[model['fields']['kana'],block('氏名',8),model['fields']['name'],model['fields']['birth']])
    if d.get('photo'):
        right.text='';p=right.paragraphs[0];p.paragraph_format.line_spacing=Pt(114);p.add_run().add_picture(io.BytesIO(base64.b64decode(d['photo'].split(',',1)[1])),width=Mm(30),height=Mm(40))
    else:write(right,[block('写真を貼る位置',8),block('縦40mm × 横30mm',8)])
    for cell,l,r in [(first[2],'address','contact'),(first[3],'alternate','alternatePhone')]:nested(cell,[[model['fields'][l],model['fields'][r]]],[W-144,144],[46])
    chronology_cell(first[5],model['jobs'][0],TITLE)
    paragraph=doc.add_paragraph();paragraph.paragraph_format.line_spacing=Pt(1);paragraph.add_run().add_break(WD_BREAK.PAGE)
    heights=model['page2']
    second=outer([heights['chronology'],12,heights['qualifications'],12,heights['motivation'],12,heights['preferences']])
    chronology_cell(second[0],model['jobs'][1],TITLE);chronology_cell(second[2],model['qualifications'],'免許・資格')
    for cell,key,label,labelh in [(second[4],'motivation',MOTIVATION,22),(second[6],'preferences',PREFERENCES,30)]:
        labelblock=fit(label,W-12,labelh,label,8,8);nested(cell,[[labelblock],[model['fields'][key]]],[W],[labelh,heights[key]-labelh-2])
    doc.add_paragraph().paragraph_format.line_spacing=Pt(1)
    doc.save(out)
