"""Local parsing and Japanese document export. Input documents are data, never instructions."""
import sys,json,os,io,base64,zipfile,subprocess,shutil,tempfile,re
from xml.etree import ElementTree as ET
from pathlib import Path
import fitz
from PIL import Image, ImageOps, ImageSequence

def ocr_image(raw):
    image=ImageOps.exif_transpose(Image.open(io.BytesIO(raw))).convert('RGB')
    image.thumbnail((2800,3600))
    buffer=io.BytesIO(); image.save(buffer,format='PNG')
    if sys.platform=='darwin':
        try:
            with tempfile.TemporaryDirectory(prefix='ma-ocr-') as folder:
                source=Path(folder)/'input.png';source.write_bytes(buffer.getvalue())
                result=subprocess.run(['/usr/bin/osascript','-l','JavaScript',str(Path(__file__).with_name('mac_ocr.js')),str(source)],capture_output=True,timeout=25)
                if result.returncode==0:
                    value=json.loads(result.stdout.decode()).get('text')
                    if isinstance(value,str) and value.strip():return value.strip()
        except (ValueError,KeyError,OSError,subprocess.TimeoutExpired):pass
        if not shutil.which('tesseract'):raise ValueError('Macの画像読取を完了できませんでした。画像を再度選択するかGeminiで確認してください。')
    if not shutil.which('tesseract'):raise ValueError('画像のローカルOCRを利用できません。Geminiによる画像読取が必要です。')
    result=subprocess.run(['tesseract','stdin','stdout','--tessdata-dir',str(Path(__file__).parent/'ocr'),'-l','jpn+eng','--psm','3'],input=buffer.getvalue(),capture_output=True,timeout=12,env={**os.environ,'OMP_THREAD_LIMIT':'2'})
    if result.returncode: raise ValueError('OCRで画像を読めませんでした。画像の向き・解像度を確認してください。')
    return result.stdout.decode('utf-8').strip()

def extract(path, original):
    ext=Path(original).suffix.lower(); text=''; photos=[]; pages=[]; warnings=[]; text_ready=False; ocr_ready=False; page_count=0; missing_parts=False; unreadable_media=False
    if ext in ['.docx','.xlsx','.xlsm','.pptx']:
        with zipfile.ZipFile(path) as z:
            if len(z.infolist())>10000 or sum(i.file_size for i in z.infolist())>200*1024*1024:
                raise ValueError('展開後の書類サイズが大きすぎます。資料を分割してください。')
    def save_photo(raw):
        nonlocal missing_parts,unreadable_media
        office=ext in ['.docx','.xlsx','.xlsm','.pptx']
        try:
            with Image.open(io.BytesIO(raw)) as opened:
                im=ImageOps.exif_transpose(opened).convert('RGB');w,h=im.size
            if office:
                readable=im.copy();readable.thumbnail((2400,3200))
                b=io.BytesIO();readable.save(b,format='JPEG',quality=92)
                pages.append('data:image/jpeg;base64,'+base64.b64encode(b.getvalue()).decode())
            if w<80 or h<80:return
            im.thumbnail((900,1200));b=io.BytesIO();im.save(b,format='JPEG',quality=90)
            photos.append({'data':'data:image/jpeg;base64,'+base64.b64encode(b.getvalue()).decode(),'width':w,'height':h})
        except Exception:
            if office:missing_parts=True;unreadable_media=True
    if ext=='.pdf':
        doc=fitz.open(path)
        if doc.is_encrypted: raise ValueError('パスワード付きPDFは解除してからアップロードしてください。')
        if len(doc)>40: raise ValueError('PDFは40ページ以内に分割してください。')
        page_count=len(doc)
        seen=set(); native_complete=True; all_pages_read=True
        for page in doc:
            page_text=page.get_text()
            page_native=len(page_text.strip())>=80 and '\ufffd' not in page_text
            if not page_native:native_complete=False
            for image_info in page.get_images(full=True):
                try:
                    for rect in page.get_image_rects(image_info[0]):
                        fraction=rect.get_area()/page.rect.get_area()
                        portrait=(.55<rect.width/max(rect.height,1)<.95 and rect.x0>page.rect.width*.65 and rect.y0<page.rect.height*.25 and fraction<.08)
                        if fraction>.015 and not portrait: native_complete=False;page_native=False
                except Exception: native_complete=False;page_native=False
            page_read=page_native
            if not page_native:
                try:
                    ocr_text=ocr_image(page.get_pixmap(matrix=fitz.Matrix(2.5,2.5)).tobytes('png'))
                    if ocr_text.strip():
                        page_read=True
                        page_text=(page_text+'\n【画像OCR】\n'+ocr_text) if page_text.strip() else ocr_text
                    warnings.append(f'{page.number+1}ページ目を無料OCRで読み取りました。誤字・項目の対応を確認してください。')
                except (ValueError,subprocess.TimeoutExpired) as e: warnings.append(f'{page.number+1}ページ目のOCRが完了しませんでした。手入力で確認してください。')
            all_pages_read=all_pages_read and page_read
            text+=page_text+'\n'
            for img in page.get_images(full=True):
                if img[0] not in seen:
                    seen.add(img[0]); save_photo(doc.extract_image(img[0])['image'])
            if len(pages)<8:
                scale=min(2.5,2400/max(page.rect.width,page.rect.height))
                pix=page.get_pixmap(matrix=fitz.Matrix(scale,scale),alpha=False)
                pages.append('data:image/jpeg;base64,'+base64.b64encode(pix.tobytes('jpeg',jpg_quality=90)).decode())
        text_ready=native_complete and len(text)<=100000
        ocr_ready=not text_ready and all_pages_read and bool(text.strip())
        if not text.strip(): warnings.append('テキストを読み取れませんでした。原本を参照して手入力してください。')
        if len(doc)>8: warnings.append('AIへ送る画像は先頭8ページです。テキストは全ページから抽出し、必要なページにはOCRを行っています。')
    elif ext=='.docx':
        # Word's paragraph API omits floating text boxes (often contact details).
        # Read all visible runs in document order, including nested tables and boxes.
        ns='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        with zipfile.ZipFile(path) as z:
            parts=[]
            for n in z.namelist():
                if n=='word/document.xml' or (n.startswith(('word/header','word/footer')) and n.endswith('.xml')):
                    chunks=[]
                    for event,node in ET.iterparse(io.BytesIO(z.read(n)),events=('start','end')):
                        if event=='end' and node.tag==ns+'t': chunks.append(node.text or '')
                        elif event=='end' and node.tag in [ns+'tab',ns+'br',ns+'cr']: chunks.append('\t' if node.tag==ns+'tab' else '\n')
                        elif event=='end' and node.tag in [ns+'p',ns+'tc']: chunks.append('\n')
                    parts.append(''.join(chunks))
                if n.startswith('word/media/'): save_photo(z.read(n))
            text='\n'.join(parts)
    elif ext in ['.xlsx','.xlsm']:
        import openpyxl
        w=openpyxl.load_workbook(io.BytesIO(Path(path).read_bytes()),read_only=False,data_only=True)
        formulas=openpyxl.load_workbook(io.BytesIO(Path(path).read_bytes()),read_only=True,data_only=False)
        count=0
        for sheet in w:
            text+='\n'+sheet.title+'\n'
            for row,formula_row in zip(sheet.iter_rows(),formulas[sheet.title].iter_rows()):
                count+=1
                if count>15000: raise ValueError('Excelは15,000行以内に分割してください。')
                values=[]
                for cell,formula in zip(row,formula_row):
                    if formula.data_type=='f' and cell.value is None:
                        missing_parts=True
                        values.append('[計算結果未保存:'+str(formula.value)+']')
                    elif cell.value is not None:values.append(str(cell.value))
                text+=' | '.join(values)+'\n'
            for im in sheet._images: save_photo(im._data())
        w.close();formulas.close()
        if missing_parts:warnings.append('Excelの数式に保存済みの計算結果がないセルがあります。数式は保持し、計算結果を推測していません。')
    elif ext=='.pptx':
        from pptx import Presentation
        deck=Presentation(path);page_count=len(deck.slides)
        def slide_text(shapes):
            nonlocal text
            for s in shapes:
                if s.has_text_frame:text+=s.text+'\n'
                if s.has_table:
                    for row in s.table.rows:text+=' | '.join(cell.text for cell in row.cells)+'\n'
                if s.shape_type==6:slide_text(s.shapes)
                if s.shape_type==13:save_photo(s.image.blob)
        for slide in deck.slides:
            slide_text(slide.shapes)
            if slide.has_notes_slide:text+=slide.notes_slide.notes_text_frame.text+'\n'
    elif ext in ['.txt','.csv','.md','.json','.rtf']:
        raw=Path(path).read_bytes()
        encodings=['utf-32'] if raw.startswith((b'\xff\xfe\x00\x00',b'\x00\x00\xfe\xff')) else ['utf-16'] if raw.startswith((b'\xff\xfe',b'\xfe\xff')) else ['utf-8-sig','cp932']
        if ext=='.rtf':
            codepage=re.search(rb'\\ansicpg(\d+)',raw)
            encodings=['cp'+codepage.group(1).decode()] if codepage else ['cp1252']
        for enc in encodings:
            try: text=raw.decode(enc); break
            except (UnicodeError,LookupError): pass
        if not text: raise ValueError('文字コードを読めません。UTF-8で保存してください。')
        if ext=='.rtf':
            from striprtf.striprtf import rtf_to_text
            text=rtf_to_text(text)
    elif ext in ['.jpg','.jpeg','.png','.webp','.tif','.tiff','.bmp','.heic']:
        try:image_file=Image.open(path)
        except (OSError,ValueError):raise ValueError('画像を開けません。JPEGまたはPNGに変換してください。')
        with image_file as frames:
            page_count=getattr(frames,'n_frames',1)
            if page_count>8:raise ValueError('画像は一度に8ページ以内で読み取ってください。内容は省略していません。')
            pages=[];texts=[];ocr_ready=True
            for frame in ImageSequence.Iterator(frames):
                image=ImageOps.exif_transpose(frame.copy()).convert('RGB');image.thumbnail((2400,3200))
                buffer=io.BytesIO();image.save(buffer,format='JPEG',quality=92);raw=buffer.getvalue()
                pages.append('data:image/jpeg;base64,'+base64.b64encode(raw).decode());photos=[]
                try:
                    value=ocr_image(raw);texts.append(value);ocr_ready=ocr_ready and bool(value.strip())
                except subprocess.TimeoutExpired:ocr_ready=False;warnings.append('画像読取の処理時間を超えました。Geminiでの確認が必要です。')
                except ValueError as error:ocr_ready=False;warnings.append(str(error))
            text='\n'.join(texts)
        warnings.append(('無料OCRで読み取りました。' if text else 'ローカルOCRでは読み取れませんでした。Gemini設定時は画像を送信して読み取れます。')+'原本と照合してください。顔写真はプレビューで切り抜いて再利用できます。')
    else:
        raise ValueError('この形式は直接読取に対応していません。PDF・DOCX・XLSX・PPTX・画像・TXT・CSVに変換してください（旧DOC/XLSは変換が必要です）。')
    if unreadable_media:raise ValueError('資料内に読み取れない画像形式があります。内容は省略せず停止しました。元のアプリからPDFとして保存したファイルを使用してください。')
    if ext in ['.docx','.xlsx','.xlsm','.pptx'] and len(pages)>8:raise ValueError('資料内の画像が8点を超えています。内容は省略していません。')
    if ext in ['.docx','.xlsx','.xlsm','.pptx'] and pages:
        ocr_ready=True
        for data in pages:
            try:
                value=ocr_image(base64.b64decode(data.split(',',1)[1]));ocr_ready=ocr_ready and bool(value.strip())
                if value.strip():text+='\n【資料内画像のOCR】\n'+value
            except (ValueError,subprocess.TimeoutExpired):ocr_ready=False
        warnings.append('資料内の画像にもOCRを行いました。読取結果を原本と照合してください。')
    if ext in ['.docx','.xlsx','.xlsm','.pptx','.txt','.csv','.md','.json','.rtf']:
        text_ready=bool(text.strip()) and not pages and not missing_parts
    if len(text)>100000:raise ValueError('資料が100,000文字を超えています。読取内容は省略していません。')
    ocr_ready=ocr_ready and not missing_parts
    photos.sort(key=lambda p: (0 if .55<p['width']/p['height']<.95 else 1,-p['width']*p['height']))
    return {'text':text, 'photos':photos[:10], 'pages':pages,'warnings':warnings,'textReady':text_ready,'ocrTextReady':ocr_ready,'pageCount':page_count}

def sections(d):
    rows=lambda k: '\n'.join('　'.join(str(r.get(c,'')) for c in ['date','end','name','detail']) for r in d.get(k,[]))
    if d.get('type')=='shokumukeirekisho':
        return [('職務要約',d.get('summary','')),('活かせる経験・スキル',d.get('skills','')),('職務経歴',rows('experience')),('PC・技術スキル',d.get('techSkills','')),('自己PR・補足事項',d.get('additional',''))]
    return [('学歴',rows('education')),('職歴',rows('experience')),('免許・資格',rows('licenses')),('志望の動機',d.get('motivation','')),('特技・自己PR',d.get('strengths','')),('本人希望記入欄',d.get('preferences',''))]

def export(d,fmt,out):
    from export_documents import render_document
    return render_document(d,fmt,out)

if __name__=='__main__':
    try:
        if sys.argv[1]=='check':
            from resume_layout import layout
            layout(json.load(sys.stdin));print(json.dumps({'pages':2,'fits':True}))
        elif sys.argv[1]=='extract': print(json.dumps(extract(sys.argv[2],sys.argv[3]),ensure_ascii=False))
        else:
            export(json.load(sys.stdin),sys.argv[2],sys.argv[3]); print(json.dumps({'ok':True}))
    except Exception as e:
        print(json.dumps({'error':str(e)},ensure_ascii=False)); sys.exit(1)
